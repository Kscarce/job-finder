"""
Reads your resumes from the resume/ folder and pulls out the text.
Works with .docx, .pdf and .txt files.

Run it on its own to check the text came out right:
    python jobfinder/resume_reader.py
"""

from pathlib import Path

import config

RESUME_TYPES = (".docx", ".pdf", ".txt")


def read_text(path):
    """Return the plain text of one resume file."""
    path = Path(path)
    if path.suffix.lower() == ".docx":
        import docx  # the python-docx package
        document = docx.Document(path)
        lines = [p.text for p in document.paragraphs]
        # Text inside tables (some resume templates use them for layout).
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    return path.read_text(encoding="utf-8", errors="ignore")


def load_resumes():
    """
    Returns {label: text} for every resume in the resume/ folder.
    The label comes from the file name: resume_ml_research.docx -> "ml research".
    """
    resumes = {}
    config.RESUME_DIR.mkdir(exist_ok=True)
    for path in sorted(config.RESUME_DIR.iterdir()):
        if path.name.lower() == "readme.txt":
            continue
        if path.suffix.lower() in RESUME_TYPES and not path.name.startswith("~$"):
            label = path.stem.lower().removeprefix("resume").strip("_- ").replace("_", " ")
            resumes[label or path.stem] = read_text(path)
    return resumes


SECTION_HEADINGS = {
    "education", "projects", "technical skills", "skills", "experience",
    "work experience", "leadership & involvement", "leadership", "summary",
}


def split_sections(text):
    """
    Splits a resume into {heading: text}. Everything before the first heading
    (name, contact line, summary paragraph) goes under "summary".
    """
    sections, current = {"summary": []}, "summary"
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.lower() in SECTION_HEADINGS:
            current = line.lower()
            sections.setdefault(current, [])
        else:
            sections[current].append(line)
    return {k: "\n".join(v) for k, v in sections.items() if v}


def background_pieces(resumes):
    """
    Builds the separate pieces of your background that each job is compared to.
    Lines from all resumes are combined (without repeats), so something on only
    one resume still counts.
    Returns {piece name: text}.
    """
    pieces = {"skills": [], "degree and coursework": [], "dissertation": [], "leadership": []}

    def add(piece, text):
        for line in text.splitlines():
            if line not in pieces[piece]:
                pieces[piece].append(line)

    for text in resumes.values():
        s = split_sections(text)
        # The summary paragraph is the last line before the first heading.
        add("skills", s.get("summary", "").splitlines()[-1] if s.get("summary") else "")
        add("skills", s.get("technical skills", "") or s.get("skills", ""))
        add("degree and coursework", s.get("education", ""))
        add("leadership", s.get("leadership & involvement", "") or s.get("leadership", ""))
        add("leadership", s.get("experience", "") or s.get("work experience", ""))

        # Projects: the dissertation is its own piece, everything else is coursework.
        # A project is a title line followed by its description line.
        lines = s.get("projects", "").splitlines()
        for title, description in zip(lines[0::2], lines[1::2]):
            target = "dissertation" if "dissertation" in title.lower() else "degree and coursework"
            add(target, title + "\n" + description)

    return {name: "\n".join(l for l in lines if l) for name, lines in pieces.items()}


if __name__ == "__main__":
    resumes = load_resumes()
    for label, text in resumes.items():
        print(f"===== {label} ({len(text.split())} words) =====")
        print(text)
        print()
    print("===== background pieces each job is compared to =====")
    for name, text in background_pieces(resumes).items():
        print(f"\n--- {name} ---\n{text}")
